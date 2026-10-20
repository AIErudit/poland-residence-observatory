"""Independently reconcile every remaining curated value to acquired sources.

Reads the external acquisition cache; writes only an aggregate verification result.
No raw PDFs, workbooks, letters or articles are copied into the repository.
"""
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone
from urllib.parse import urlparse, parse_qs
import contextlib
import hashlib
import io
import json
import math
import re
import runpy
import unicodedata
import openpyxl
import pypdf
import xlrd
from bs4 import BeautifulSoup

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT.parent/'source-cache'/'observatory'/'blobs'
with contextlib.redirect_stdout(io.StringIO()):
    core = runpy.run_path(str(ROOT/'verification/round2/recheck_core.py'))
data = core['data']
obs = data['observations']
sources = {s['id']:s for s in data['sources']}
verified = set(core['verified_observation_ids'])
core_count = len(verified)
issues = list(core['issues'])
counts = Counter()
read_sources = {}
geos = core['geos']

def require(condition, message):
    if not condition:
        issues.append(message)

def raw(source_id):
    s = sources[source_id]
    h = s.get('content_sha256')
    if not h:
        raise ValueError(f'No acquisition hash for {source_id}')
    content = (CACHE/(h+'.bin')).read_bytes()
    if hashlib.sha256(content).hexdigest() != h:
        raise ValueError(f'Cached content hash differs for {source_id}')
    read_sources[source_id] = {'url':s['url'],'sha256':h,'access':s['access']}
    return content

def compare(o, expected):
    equal = o['value'] == expected if expected is None or isinstance(expected,int) else (
        isinstance(o['value'],(int,float)) and math.isclose(o['value'],expected,rel_tol=0,abs_tol=0.000001))
    require(equal, f"Value {o['id']}: {o['series_id']}/{o['indicator_id']}: {o['value']} != {expected}")
    if equal and o['id'] not in verified:
        verified.add(o['id'])
        counts[o['series_id']] += 1

def selected(series_id, **fields):
    return [o for o in obs if o['series_id']==series_id and all(o.get(k)==v for k,v in fields.items())]

def norm(text):
    return ''.join(c for c in unicodedata.normalize('NFD',str(text).lower()) if not unicodedata.combining(c)).replace('ł','l')

# Remaining NSA fields: read actual headers of each original XLS layout.
for sid in sorted({o['source_id'] for o in selected('nsa_wsa_6272_annual')}):
    book = xlrd.open_workbook(file_contents=raw(sid))
    header = book.sheet_by_name('WSA SAB Ogółem')
    years = {int(y) for i in range(3) for v in header.row_values(i)
             for y in re.findall(r'202[0-5]',str(v))}
    require(len(years)==1, f'NSA ambiguous report year {sid}')
    if len(years)!=1:
        continue
    year = str(next(iter(years)))
    extracted = {}
    for sh in book.sheets():
        if sh.name=='WSA SAB Ogółem':
            continue
        matches=[i for i in range(sh.nrows) if any(v==6272 for v in sh.row_values(i))]
        require(len(matches)==1, f'NSA row6272 cardinality {year}/{sh.name}')
        if len(matches)!=1:
            continue
        r = sh.row_values(matches[0])
        cells=[(j,norm(v)) for i in range(10) for j,v in enumerate(sh.row_values(i)) if str(v).strip()]
        def unique_col(predicate, label):
            cols={j for j,v in cells if predicate(v)}
            if len(cols)!=1:
                raise ValueError(f'NSA header {label} ambiguous {year}/{sh.name}: {cols}')
            return next(iter(cols))
        start=unique_col(lambda v:'pozostalo z poprzedniego' in v,'pending_start')
        end=unique_col(lambda v:'pozostalo na nastepny' in v,'pending_end')
        closed=unique_col(lambda v:v.startswith('zamkn') or v.startswith('zamnk'),'closed')
        combined=unique_col(lambda v:'lacznie' in v and 'kol' in v and '17' in v and 'suma rubryk' not in v,'combined')
        resolved=unique_col(lambda v:'lacznie' in v and 'suma rubryk' in v,'resolved')
        def n(col):
            return int(r[col] or 0)
        outcomes={}
        for name,term in [('up','uwzgledniono'),('down','oddalono'),('rejected','odrzucono')]:
            cols={j for j,v in cells if term in v}
            outcomes[name]=sum(n(j) for j in cols)
        expected={
            'court_pending_start':n(start),'court_pending_end':n(end),
            'court_closed':n(closed),'court_resolved_including_closed':n(combined),
            'court_other_resolved':n(resolved)-sum(outcomes.values())}
        expected.update({'court_resolved':n(resolved),'court_inflow':n(4),
                         'court_upheld':outcomes['up'],'court_dismissed':outcomes['down'],
                         'court_rejected':outcomes['rejected']})
        require(expected['court_pending_start']+expected['court_inflow']-expected['court_resolved_including_closed']==expected['court_pending_end'],
                f'NSA movement identity {year}/{sh.name}')
        require(expected['court_resolved']+expected['court_closed']==expected['court_resolved_including_closed'],
                f'NSA closure identity {year}/{sh.name}')
        gid='PL' if sh.name=='WSA SAB Symbole' else 'WSA-'+sh.name
        extracted[gid]=expected
        for o in selected('nsa_wsa_6272_annual',source_id=sid,geography_id=gid):
            if o['indicator_id'] in expected:
                compare(o,expected[o['indicator_id']])
                if o['indicator_id']=='court_pending_start':
                    require(o['as_of_date']==year+'-01-01',f"Pending-start date {o['id']}")
    for key in ['court_pending_start','court_pending_end','court_closed','court_resolved_including_closed','court_other_resolved']:
        require(sum(v[key] for gid,v in extracted.items() if gid!='PL')==extracted['PL'][key],f'NSA16-court sum {year}/{key}')

# UdSC Blue Card table, complete public scan page4: all64 values visually read.
blue=[
 [277,332,158,236],[407,401,435,366],[166,358,287,412],[271,296,512,376],
 [290,251,260,307],[130,123,156,211],[105,155,194,313],[419,594,489,314],
 [121,191,247,329],[124,147,239,457],[149,156,353,412],[291,386,582,568],
 [95,248,289,438],[99,139,174,264],[210,240,360,350],[231,238,358,366]]
for o in selected('udsc_20260828_blue_card_mean_days'):
    content=raw(o['source_id'])
    year=int(o['reference_period'][:4])
    compare(o,blue[geos.index(o['geography_id'])][year-2023])
    require(o['permit_type']=='blue_card',f"Blue Card scope {o['id']}")
    require(year<2026 or o['as_of_date']=='2026-08-12',f"Blue Card cut-off {o['id']}")

# Original two-page warehouse PDFs: names/counts/means on page1, maxima on page2.
for sid in sorted({o['source_id'] for o in selected('warehouse_all_outcomes')}):
    pdf=pypdf.PdfReader(io.BytesIO(raw(sid)))
    first=pdf.pages[0].extract_text()
    records=re.findall(r'(WOJEWODA[^\n]+)\n(\d+)\n(\d+)',first)
    maxima=[int(v) for v in re.findall(r'^\s*(\d+)\s*$',pdf.pages[1].extract_text(),re.M)]
    require(len(records)==len(maxima)==16,f'Warehouse row cardinality {sid}')
    if len(records)!=16 or len(maxima)!=16:
        continue
    for o in selected('warehouse_all_outcomes',source_id=sid):
        i=geos.index(o['geography_id'])
        values={'concluded_resolutions':int(records[i][1]),'processing_duration_mean_days':int(records[i][2]),'processing_duration_max_days':maxima[i]}
        compare(o,values[o['indicator_id']])
        if o['reference_period']=='2026':
            require(o['as_of_date'] is None and o['quality_status']=='reviewed_primary_period_ambiguous',f"Warehouse2026 uncertainty {o['id']}")

# Łódź workbook: infer category/year from the actual A1/title, not observation locators.
lodz_series=['lodz_warehouse_positive_decisions','lodz_warehouse_all_outcomes']
lodz_sids={o['source_id'] for o in obs if o['series_id'] in lodz_series}
require(len(lodz_sids)==1,'Multiple Lodz workbooks')
for sid in lodz_sids:
    book=openpyxl.load_workbook(io.BytesIO(raw(sid)),data_only=True)
    require(len(book.sheetnames)==24,'Expected24 Lodz sheets')
    for sh in book:
        title=sh.title
        definition=str(sh['A1'].value or '')
        context=title+' '+definition
        yrs=re.findall(r'202[456]',context)
        require(bool(yrs) and len(set(yrs))==1,f'Workbook year {title}')
        if not yrs:
            continue
        year=yrs[0]
        positive='decyzji pozytywnej' in definition
        if positive:
            arts=re.findall(r'art\.?\s*(114|127|142|144|158)',definition)
            require(len(set(arts))==1,f'Workbook article {title}')
            if not arts:
                continue
            permits={'114':'temporary_work_art114','127':'blue_card_art127','142':'temporary_business_art142','144':'temporary_study_art144','158':'temporary_family_art158'}
            permit=permits[arts[0]]
            series_id='lodz_warehouse_positive_decisions'
            require('Usunięto duplikaty' in str(sh['G3'].value),f'Validation note {title}')
        else:
            require('negatywne' in definition and 'umorzenia' in definition and 'bez rozpoznania' in definition,f'All-outcome definition {title}')
            permit='permanent_residence' if 'stały' in title else 'eu_long_term_resident' if 'Rezydent' in title else 'temporary_residence'
            series_id='lodz_warehouse_all_outcomes'
        rows=[r for r in range(2,6) if sh.cell(r,1).value=='WOJEWODA ŁÓDZKI']
        require(len(rows)==1,f'Lodz data-row cardinality {title}')
        if len(rows)!=1:
            continue
        rr=rows[0]
        values=dict(zip(['concluded_resolutions','processing_duration_mean_days','processing_duration_max_days'],[sh.cell(rr,col).value for col in [2,3,4]]))
        items=[o for o in selected(series_id,source_id=sid,permit_type=permit) if o['reference_period'][:4]==year]
        require(len(items)==3,f'Lodz observation coverage {title}')
        for o in items:
            compare(o,values[o['indicator_id']])
            require(o['stage']==('concluded_positive' if positive else 'concluded'),f"Lodz cohort {o['id']}")
            if year=='2026':
                require('31.08' in context and o['as_of_date']=='2026-08-31',f"Lodz2026 cut-off {o['id']}")

# Direct cached UdSC API responses: all16 national totals; no request is sent.
for o in [o for o in obs if o['series_id'].startswith('udsc_portal_')]:
    s=sources[o['source_id']]
    parsed=urlparse(s['url']);params=parse_qs(parsed.query)
    rows=json.loads(raw(o['source_id']))
    require(params.get('groupBy')==['country'],f"Portal grouping {o['id']}")
    require(len({r.get('country') for r in rows})==len(rows),f"Portal duplicate group {o['id']}")
    compare(o,sum(int(r['total']) for r in rows))
    require(params.get('year')==[o['reference_period'][:4]],f"Portal year {o['id']}")
    case='3' if o['permit_type']=='eu_long_term_resident' else '1'
    require(params.get('caseType')==[case],f"Portal caseType {o['id']}")
    measure='applications' if o['indicator_id']=='applications_received' else 'decisions'
    require('/'+measure+'/poland' in parsed.path,f"Portal measure {o['id']}")

# MSWiA12584 table: exactly two populations; accompanying mean has an unknown cohort.
art=selected('mswia_12584_pending_count')+selected('mswia_12584_reported_mean_days')
for sid in {o['source_id'] for o in art}:
    soup=BeautifulSoup(raw(sid),'html.parser')
    table=next(t for t in soup.find_all('table') if '463 487' in t.get_text(' ',strip=True))
    rows={}
    for tr in table.find_all('tr'):
        cells=[c.get_text(' ',strip=True) for c in tr.find_all(['td','th'])]
        if cells and ('Wojewoda Opolski' in cells[0] or 'Wszyscy wojewodowie' in cells[0]):
            key='PL-OP' if 'Opolski' in cells[0] else 'PL'
            rows[key]=[int(re.sub(r'\s','',c)) for c in cells[1:3]]
    for o in [o for o in art if o['source_id']==sid]:
        compare(o,rows[o['geography_id']][0 if o['indicator_id']=='pending_cases' else 1])
        require(o['as_of_date']=='2025-10-13',f"MSWiA12584 date {o['id']}")
        if o['indicator_id']!='pending_cases':
            require(o['definition_id']=='mswia12584_mean_ambiguous',f"MSWiA12584 mean scope {o['id']}")

remaining=[{k:o[k] for k in ['id','series_id','indicator_id','source_id']} for o in obs if o['id'] not in verified]
result={
 'checked_at':datetime.now(timezone.utc).isoformat(),
 'status':'all_numeric_values_reconciled' if not issues and not remaining else 'blocking_or_unverified_rows',
 'dataset_sha256':hashlib.sha256((ROOT/'data/curated/dataset.json').read_bytes()).hexdigest(),
 'core_previously_reconciled':core_count,'additional_reconciled':len(verified)-core_count,
 'total_numeric_observations_reconciled':len(verified),'total_observations':len(obs),
 'additional_counts_by_series':dict(counts),'source_snapshots_read':read_sources,
 'issues':issues,'unverified_rows':remaining,
 'limits':['Numeric agreement with acquired source snapshots does not authenticate signed originals or validate administrative source systems.',
           'Redacted UdSC regional labels remain reconstructed; all-copy provenance qualifiers remain applicable.',
           'Warehouse2026 values match the table but exact operational cut-off is unknown; as_of_date=null and period_ambiguous are required, no annual projection permitted.',
           'Reported maxima are database statistics, not independently authenticated individual histories.',
           'All1483 catalogue URLs, all6625 upstream CSV records and all research claims are not independently fact-checked by this numeric reconciliation.']}
(ROOT/'verification/round2/numeric-recheck.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps({k:v for k,v in result.items() if k not in ['source_snapshots_read','unverified_rows']},ensure_ascii=False))
print('unverified_row_count',len(remaining))
